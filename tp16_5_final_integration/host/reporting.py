from __future__ import annotations

import csv
import html
import json
from pathlib import Path
from typing import Any


def _safe(value: Any) -> str:
    return html.escape(str(value))


def save_json(session: dict, path: Path) -> None:
    path.write_text(
        json.dumps(session, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def save_csv(session: dict, path: Path) -> None:
    samples = session.get("samples", [])
    fieldnames = [
        "elapsed_s",
        "timestamp",
        "state",
        "auto",
        "sensor_valid",
        "distance_valid",
        "obstacle",
        "temperature_c",
        "humidity_pct",
        "pressure_hpa",
        "distance_cm",
        "limit_open",
        "limit_closed",
        "contradiction",
        "ref_valid",
        "ref_temperature_c",
        "ref_humidity_pct",
        "ref_pressure_hpa",
        "delta_temperature_c",
        "delta_humidity_pct",
        "delta_pressure_hpa",
    ]

    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for sample in samples:
            writer.writerow({key: sample.get(key, "") for key in fieldnames})


def save_html(session: dict, path: Path) -> None:
    summary = session.get("summary", {})
    checks = session.get("checks", [])
    samples = session.get("samples", [])
    events = session.get("events", [])

    check_rows = "".join(
        f"<tr><td>{_safe(c['name'])}</td><td>{_safe(c['result'])}</td><td>{_safe(c.get('details', ''))}</td></tr>"
        for c in checks
    )

    event_rows = "".join(
        f"<tr><td>{_safe(e['elapsed_s'])}</td><td>{_safe(e['event'])}</td><td>{_safe(e.get('details', ''))}</td></tr>"
        for e in events
    )

    sample_rows = "".join(
        "<tr>"
        f"<td>{_safe(s.get('elapsed_s', ''))}</td>"
        f"<td>{_safe(s.get('state', ''))}</td>"
        f"<td>{_safe(s.get('temperature_c', ''))}</td>"
        f"<td>{_safe(s.get('humidity_pct', ''))}</td>"
        f"<td>{_safe(s.get('distance_cm', ''))}</td>"
        f"<td>{_safe(s.get('obstacle', ''))}</td>"
        f"<td>{_safe(s.get('ref_temperature_c', ''))}</td>"
        f"<td>{_safe(s.get('delta_temperature_c', ''))}</td>"
        "</tr>"
        for s in samples
    )

    doc = f"""<!doctype html>
<html lang="fr">
<head>
<meta charset="utf-8">
<title>TP16.5 - Demo Report</title>
<style>
body {{ font-family: Arial, sans-serif; margin: 32px; color: #202124; }}
h1, h2 {{ margin-bottom: 8px; }}
table {{ border-collapse: collapse; width: 100%; margin: 12px 0 28px; }}
th, td {{ border: 1px solid #bbb; padding: 7px; text-align: left; }}
th {{ background: #f2f2f2; }}
code {{ background: #f4f4f4; padding: 2px 4px; }}
</style>
</head>
<body>
<h1>TP16.5 - Final Integrated Demo</h1>
<p><strong>Session:</strong> {_safe(session.get('started_at', ''))}</p>
<p><strong>BENCH:</strong> {_safe(session.get('versions', {}).get('bench', 'unknown'))} &nbsp; <strong>DUT:</strong> {_safe(session.get('versions', {}).get('dut', 'unknown'))}</p>
<p><strong>Overall:</strong> {_safe(summary.get('overall', 'RECORDED'))}</p>

<h2>Checks</h2>
<table><tr><th>Check</th><th>Result</th><th>Details</th></tr>{check_rows}</table>

<h2>Events</h2>
<table><tr><th>Elapsed (s)</th><th>Event</th><th>Details</th></tr>{event_rows}</table>

<h2>Samples</h2>
<table>
<tr><th>s</th><th>State</th><th>T DUT</th><th>H DUT</th><th>Distance</th><th>Obstacle</th><th>T REF</th><th>dT</th></tr>
{sample_rows}
</table>
</body>
</html>"""
    path.write_text(doc, encoding="utf-8")


def save_xlsx(session: dict, path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for XLSX export. Run: pip install -r host/requirements.txt"
        ) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    summary_rows = [
        ("Started", session.get("started_at", "")),
        ("Finished", session.get("finished_at", "")),
        ("BENCH version", session.get("versions", {}).get("bench", "unknown")),
        ("DUT version", session.get("versions", {}).get("dut", "unknown")),
        ("Overall", session.get("summary", {}).get("overall", "RECORDED")),
        ("States seen", ", ".join(session.get("summary", {}).get("states_seen", []))),
        ("Samples", len(session.get("samples", []))),
    ]

    for row in summary_rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True)

    checks_ws = wb.create_sheet("Checks")
    checks_ws.append(["Check", "Result", "Details"])
    for cell in checks_ws[1]:
        cell.font = Font(bold=True)
    for check in session.get("checks", []):
        checks_ws.append([check.get("name"), check.get("result"), check.get("details")])

    samples_ws = wb.create_sheet("Samples")
    columns = [
        "elapsed_s", "timestamp", "state", "auto", "sensor_valid",
        "distance_valid", "obstacle", "temperature_c", "humidity_pct",
        "pressure_hpa", "distance_cm", "limit_open", "limit_closed",
        "contradiction", "ref_valid", "ref_temperature_c",
        "ref_humidity_pct", "ref_pressure_hpa", "delta_temperature_c",
        "delta_humidity_pct", "delta_pressure_hpa",
    ]
    samples_ws.append(columns)
    for cell in samples_ws[1]:
        cell.font = Font(bold=True)

    for sample in session.get("samples", []):
        samples_ws.append([sample.get(col, "") for col in columns])

    events_ws = wb.create_sheet("Events")
    events_ws.append(["elapsed_s", "event", "details"])
    for cell in events_ws[1]:
        cell.font = Font(bold=True)
    for event in session.get("events", []):
        events_ws.append([event.get("elapsed_s"), event.get("event"), event.get("details")])

    for sheet in wb.worksheets:
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            max_len = max(
                len(str(sheet.cell(row=row, column=column).value or ""))
                for row in range(1, sheet.max_row + 1)
            )
            sheet.column_dimensions[letter].width = min(max_len + 2, 35)

    wb.save(path)


def save_all(session: dict, output_dir: Path, basename: str) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)

    paths = {
        "json": output_dir / f"{basename}.json",
        "csv": output_dir / f"{basename}.csv",
        "html": output_dir / f"{basename}.html",
        "xlsx": output_dir / f"{basename}.xlsx",
    }

    save_json(session, paths["json"])
    save_csv(session, paths["csv"])
    save_html(session, paths["html"])
    save_xlsx(session, paths["xlsx"])

    return {key: str(value) for key, value in paths.items()}
